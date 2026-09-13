"""Export displayed records, preserving values and risk colors."""
from copy import copy
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side


def export_records(source, rows, destination):
    original = source['Records']
    book = Workbook()
    sheet = book.active
    sheet.title = 'Records'
    for output_row, source_row in enumerate([1, 2, 3, *rows], 1):
        for cell in original[source_row]:
            target = sheet.cell(output_row, cell.column, cell.value)
            for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
                setattr(target, attribute, copy(getattr(cell, attribute)))
            target.number_format = cell.number_format
            if output_row <= 3:
                target.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                edge = Side(style='thin', color='829AB0')
                target.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            elif original.cell(3, cell.column).value in ('dob', 'iv_date'):
                target.number_format = 'yyyy-mm-dd'
        sheet.row_dimensions[output_row].height = original.row_dimensions[source_row].height
    for merged in original.merged_cells.ranges:
        if merged.max_row <= 3:
            sheet.merge_cells(str(merged))
    for key, dimension in original.column_dimensions.items():
        sheet.column_dimensions[key].width = dimension.width
    sheet.freeze_panes = 'A4'
    sheet.auto_filter.ref = f'A3:{sheet.cell(max(3, sheet.max_row), sheet.max_column).coordinate}'
    book.save(destination)
