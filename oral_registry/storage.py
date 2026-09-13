from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

from .core import COLORS, alert_matches


def fingerprint(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class Registry:
    def __init__(self, path):
        self.path = Path(path).resolve()
        self.reload()

    def reload(self):
        self.version = fingerprint(self.path)
        self.book = openpyxl.load_workbook(self.path)
        for name in ('Records', 'Codebook', 'ตารางRiskScore', 'Entry'):
            if name not in self.book:
                raise ValueError(f'ไม่พบ Sheet {name}')
        self.columns = {}
        for c in self.book['Records'][3]:
            if c.value:
                if c.value in self.columns:
                    raise ValueError(f'ชื่อฟิลด์ซ้ำ: {c.value}')
                self.columns[c.value] = c.column
        s = self.book['Codebook']
        headers = [c.value for c in s[1]]
        self.fields = {}
        for row in s.iter_rows(min_row=2, values_only=True):
            spec = dict(zip(headers, row))
            if spec.get('column_name'):
                self.fields[spec['column_name']] = spec
        if not {'cid', 'round_id', 'riskscore', 'risklevel'} <= self.columns.keys():
            raise ValueError('Records ไม่มีคอลัมน์หลักครบ')
        if self.fields.keys() - self.columns.keys():
            raise ValueError('ฟิลด์ใน Codebook ไม่มีใน Records')
        if self.columns.keys() - self.fields.keys():
            raise ValueError('Records มีฟิลด์ที่ไม่ได้กำหนดใน Codebook กรุณาเพิ่ม mapping ก่อนใช้งาน')

    def records(self):
        s = self.book['Records']
        result = []
        for row in range(4, s.max_row + 1):
            if s.cell(row, self.columns['cid']).value is None:
                continue
            result.append((row, {key: s.cell(row, col).value for key, col in self.columns.items()}))
        return result

    def save(self, record, allow_update=False):
        lock = self.path.with_suffix(self.path.suffix + '.lock')
        try:
            fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError as exc:
            raise ValueError('มีโปรแกรมอื่นกำลังบันทึกไฟล์นี้ กรุณาลองใหม่') from exc
        temporary = None
        try:
            os.close(fd)
            if self.path.with_name('~$' + self.path.name).exists():
                raise ValueError('กรุณาปิดไฟล์ฐานข้อมูลใน Excel ก่อนบันทึก')
            if fingerprint(self.path) != self.version:
                raise ValueError('ไฟล์ถูกแก้ไขจากภายนอก กรุณาเปิดฐานข้อมูลใหม่แล้วโหลดรายการอีกครั้ง')
            key = (record['cid'], record['round_id'])
            matches = [row for row, d in self.records() if (str(d['cid']), d['round_id']) == key]
            if len(matches) > 1:
                raise ValueError('พบ CID และรอบที่ซ้ำหลายแถว ต้องแก้ข้อมูลซ้ำก่อนบันทึก')
            if matches and not allow_update:
                raise ValueError('มี CID และรอบนี้แล้ว กรุณาค้นหาและโหลดรายการเพื่อแก้ไข')
            # Read a fresh copy so a failed write never mutates the in-memory database.
            book = openpyxl.load_workbook(self.path)
            sheet = book['Records']
            occupied = [r for r in range(4, sheet.max_row + 1) if any(sheet.cell(r, c).value is not None for c in self.columns.values())]
            row = matches[0] if matches else max(occupied, default=3) + 1
            for field, col in self.columns.items():
                cell = sheet.cell(row, col)
                if row > 4:
                    cell._style = copy(sheet.cell(4, col)._style)
                value = record.get(field)
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = 's'
                if field == 'cid':
                    cell.number_format = '@'
                elif field in ('dob', 'iv_date'):
                    cell.number_format = 'yyyy-mm-dd'
                cell.fill = PatternFill('solid', fgColor=COLORS[record['risklevel']])
                font = copy(cell.font)
                font.color = '172033'
                cell.font = font
                expression = self.fields.get(field, {}).get('AlertValue')
                if field == 'r_ohplaque':
                    expression = 3  # Explicitly confirmed: plaque present receives 3 points.
                if alert_matches(value, expression):
                    cell.fill = PatternFill('solid', fgColor='FF9999')
                    font = copy(cell.font)
                    font.color = '9C0006'
                    font.bold = True
                    cell.font = font
            backups = self.path.parent / 'backups'
            backups.mkdir(exist_ok=True)
            stamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
            backup = backups / f'{self.path.stem}_{stamp}.xlsx'
            shutil.copy2(self.path, backup)
            handle, temporary = tempfile.mkstemp(suffix='.xlsx', dir=self.path.parent)
            os.close(handle)
            book.save(temporary)
            check = openpyxl.load_workbook(temporary, read_only=True)
            if check['Records'].cell(row, self.columns['cid']).value != record['cid']:
                raise ValueError('ตรวจสอบข้อมูลหลังบันทึกไม่ผ่าน')
            check.close()
            if fingerprint(self.path) != self.version:
                raise ValueError('ไฟล์เปลี่ยนระหว่างบันทึก กรุณาโหลดใหม่')
            os.replace(temporary, self.path)
            temporary = None
            self.reload()
            return row, backup
        finally:
            if temporary and Path(temporary).exists():
                Path(temporary).unlink()
            lock.unlink(missing_ok=True)

    def delete(self, cid, round_id):
        lock = self.path.with_suffix(self.path.suffix + '.lock')
        fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        os.close(fd)
        temporary = None
        try:
            if self.path.with_name('~$' + self.path.name).exists():
                raise ValueError('กรุณาปิด Excel ก่อนลบ')
            if fingerprint(self.path) != self.version:
                raise ValueError('ไฟล์เปลี่ยนจากภายนอก กรุณาเปิดรายการใหม่')
            matches = [r for r, d in self.records() if (str(d['cid']), d['round_id']) == (str(cid), round_id)]
            if len(matches) != 1:
                raise ValueError('ไม่พบรายการที่ตรงกันเพียงรายการเดียว จึงไม่ลบข้อมูล')
            book = openpyxl.load_workbook(self.path)
            book['Records'].delete_rows(matches[0])
            backups = self.path.parent / 'backups'
            backups.mkdir(exist_ok=True)
            backup = backups / (self.path.stem + '_before_delete_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f') + '.xlsx')
            shutil.copy2(self.path, backup)
            fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=self.path.parent)
            os.close(fd)
            book.save(temporary)
            check = openpyxl.load_workbook(temporary, read_only=True)
            check.close()
            if fingerprint(self.path) != self.version:
                raise ValueError('ไฟล์เปลี่ยนระหว่างลบ กรุณาเปิดรายการใหม่')
            os.replace(temporary, self.path)
            temporary = None
            self.reload()
            return backup
        finally:
            if temporary:
                Path(temporary).unlink(missing_ok=True)
            lock.unlink(missing_ok=True)
