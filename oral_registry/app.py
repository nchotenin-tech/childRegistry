from __future__ import annotations

import sys
import os
import math
import re
import json
import tempfile
import openpyxl
from datetime import date, datetime
from pathlib import Path

from PySide6.QtCore import Qt, QEvent, QSettings, QTimer
from PySide6.QtGui import QColor, QFont, QFontDatabase
from PySide6.QtWidgets import (
    QApplication, QDialog, QDialogButtonBox, QCheckBox, QFileDialog, QFormLayout, QGridLayout,
    QGroupBox, QHBoxLayout, QLabel, QLineEdit, QMainWindow, QMessageBox,
    QPushButton, QComboBox, QScrollArea, QSplitter, QTableWidget, QTableWidgetItem,
    QVBoxLayout, QWidget, QAbstractItemView,
)

from .core import COLORS, FRONT, TEETH, TOOTH_LABELS, assess, validate_record, parse_date, age_at
from .storage import Registry

ROOT = Path(__file__).resolve().parent.parent
REPEAT_FIELDS = {'round_id': 'รอบที่', 'sch': 'โรงเรียน / ศพด.', 'pv': 'จังหวัด', 'hcu': 'รพ.สต. / รพ.', 'iv_date': 'วันที่ตรวจ (DD-MM-YYYY พ.ศ.)'}


class Window(QMainWindow):
    def __init__(self, path=None):
        super().__init__()
        self.registry = None
        self.repeat_values = {}
        self.repeat_badges = {}
        self.editing = None
        self.dirty = False
        self.plaque_manual = False
        self.loading_record = False
        self.assessment_timer = QTimer(self)
        self.assessment_timer.setSingleShot(True)
        self.assessment_timer.setInterval(300)
        self.assessment_timer.timeout.connect(self.auto_assess)
        self.widgets = {}
        self.setWindowTitle('ทะเบียนสุขภาพช่องปากนักเรียน')
        self.resize(1380, 900)
        root = QWidget()
        self.setCentralWidget(root)
        layout = QVBoxLayout(root)
        title = QLabel('ทะเบียนสุขภาพช่องปากนักเรียน')
        title.setObjectName('title')
        layout.addWidget(title)
        subtitle = QLabel('ตรวจช่องปาก • สัมภาษณ์ • ประเมินความเสี่ยงฟันผุ')
        layout.addWidget(subtitle)
        bar = QHBoxLayout()
        self.file_label = QLabel('ยังไม่ได้เปิดฐานข้อมูล')
        self.file_label.setWordWrap(True)
        bar.addWidget(self.file_label, 1)
        for text, callback in [('ตั้งค่าแหล่งข้อมูล', self.open_database), ('ตั้งค่าข้อมูลใช้ซ้ำ', self.open_settings), ('ดูเด็กทั้งหมด (Records)', self.show_records), ('ติดตามผล / Dashboard', self.show_analysis)]:
            button = QPushButton(text)
            button.clicked.connect(callback)
            bar.addWidget(button)
        layout.addLayout(bar)
        split = QSplitter()
        layout.addWidget(split, 1)
        left = QWidget()
        self.search_panel = left
        toggle = QPushButton('แสดงรายการ / ค้นหา')
        toggle.setCheckable(True)
        toggle.toggled.connect(left.setVisible)
        bar.addWidget(toggle)
        left_layout = QVBoxLayout(left)
        left_layout.addWidget(QLabel('ค้นหาและเลือกรอบตรวจ'))
        self.search = QLineEdit()
        self.search.setPlaceholderText('ค้นหา CID / ชื่อเด็ก / โรงเรียน')
        self.search.textChanged.connect(self.refresh_list)
        left_layout.addWidget(self.search)
        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(['ชื่อเด็ก', 'รอบ', 'คะแนน', 'ความเสี่ยง'])
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.cellDoubleClicked.connect(self.load_selected)
        self.table.horizontalHeader().setStretchLastSection(True)
        left_layout.addWidget(self.table)
        load = QPushButton('โหลดรายการที่เลือก')
        load.clicked.connect(self.load_selected)
        left_layout.addWidget(load)
        self.count = QLabel()
        left_layout.addWidget(self.count)
        split.addWidget(left)
        right = QWidget()
        right_layout = QVBoxLayout(right)
        self.mode = QLabel('เริ่มจากตั้งค่าแหล่งข้อมูล Excel')
        self.mode.setObjectName('mode')
        self.mode.setWordWrap(True)
        self.mode.setMinimumHeight(58)
        right_layout.addWidget(self.mode)
        self.scroll = QScrollArea()
        self.scroll.setWidgetResizable(True)
        right_layout.addWidget(self.scroll, 1)
        self.result = QLabel('กรอกข้อมูลให้ครบ แล้วกดประเมินหรือบันทึก')
        self.result.setWordWrap(True)
        self.result.setObjectName('result')
        right_layout.addWidget(self.result)
        actions = QHBoxLayout()
        for text, callback in [('รายการใหม่', self.new_record), ('บันทึกข้อมูล', self.save_record)]:
            button = QPushButton(text)
            button.clicked.connect(callback)
            if text == 'บันทึกข้อมูล':
                button.setObjectName('primary')
            actions.addWidget(button)
        right_layout.addLayout(actions)
        split.addWidget(right)
        split.setSizes([300, 1080])
        left.hide()
        self.statusBar().showMessage('ข้อมูลเก็บในเครื่อง • สำรองก่อนบันทึกทุกครั้ง')
        if path:
            self.set_database(path)

    def error(self, exc):
        QMessageBox.warning(self, 'ตรวจสอบข้อมูล', str(exc))

    def discard_ok(self):
        return not self.dirty or QMessageBox.question(
            self, 'ข้อมูลยังไม่บันทึก', 'ละทิ้งการแก้ไขที่ยังไม่ได้บันทึก?',
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        ) == QMessageBox.StandardButton.Yes

    def closeEvent(self, event):
        if self.discard_ok():
            event.accept()
        else:
            event.ignore()

    def auto_assess(self):
        if not self.registry or self.loading_record or not self.dirty:
            return
        try:
            raw = {key: widget.text() for key, widget in self.widgets.items()}
            data = validate_record(raw, self.registry.fields)
            data = assess(data, self.registry.book)
        except ValueError:
            self.result.setText('ข้อมูลยังไม่ครบหรือไม่ถูกต้อง • ประเมินอัตโนมัติเมื่อกรอกครบ')
            return
        self.show_result(data)

    def changed(self, *args):
        self.dirty = True
        self.result.setText('กำลังรอข้อมูลครบเพื่อประเมินอัตโนมัติ')
        if not self.loading_record:
            self.assessment_timer.start()
        self.result.setStyleSheet('')
        if hasattr(self, 'details'):
            self.details.setRowCount(0)

    def open_database(self):
        if not self.discard_ok():
            return
        path, _ = QFileDialog.getOpenFileName(self, 'ตั้งค่าแหล่งข้อมูล Excel', str(self.registry.path if self.registry else ROOT), 'Excel (*.xlsx)')
        if path and self.set_database(path):
            QSettings('StudentOralRegistry', 'Desktop').setValue('data_source', str(self.registry.path))

    def set_database(self, path):
        try:
            registry = Registry(path)
            settings_path = registry.path.with_suffix('.settings.json')
            repeat_values = {}
            if settings_path.exists():
                payload = json.loads(settings_path.read_text(encoding='utf-8'))
                repeat_values = self.validate_settings(payload['repeat_values'])
            self.registry = registry
            self.repeat_values = repeat_values
            self.file_label.setText(str(registry.path))
            self.build_form()
            self.reset_form()
            self.refresh_list()
            return True
        except Exception as exc:
            self.error(exc)
            return False

    def show_analysis(self):
        if not self.registry:
            self.error('กรุณาตั้งค่าแหล่งข้อมูลก่อน')
            return
        from .analysis import show_dashboard
        try:
            show_dashboard(self)
        except Exception as exc:
            self.error(exc)

    def show_records(self):
        if not self.registry:
            self.error('กรุณาตั้งค่าแหล่งข้อมูลก่อน')
            return
        try:
            source = Registry(self.registry.path)
        except Exception as exc:
            self.error(exc)
            return
        dialog = QDialog(self)
        dialog.setWindowTitle('ข้อมูลเด็กทั้งหมด • Records')
        dialog.resize(1250, 760)
        layout = QVBoxLayout(dialog)
        records = source.records()
        count_label = QLabel(f'{len(records)} รอบตรวจ • {len({str(d.get("cid")) for _, d in records})} คน • อ่านจากไฟล์ล่าสุด')
        layout.addWidget(count_label)
        search = QLineEdit()
        search.setPlaceholderText('ค้นหาในทุกคอลัมน์ เช่น ชื่อ CID โรงเรียน หรือระดับความเสี่ยง')
        layout.addWidget(search)
        keys = list(source.columns)
        filters = QHBoxLayout()
        school_filter, round_filter = QComboBox(), QComboBox()
        for combo, key, title in ((school_filter, 'sch', 'ทุกโรงเรียน'), (round_filter, 'round_id', 'ทุกรอบ')):
            combo.setObjectName('filter_' + key)
            combo.addItem(title, None)
            for value in sorted({str(d.get(key, '')) for _, d in records}):
                combo.addItem(value, value)
            filters.addWidget(combo)
        reset_filters = QPushButton('ล้างตัวกรอง')
        export_button = QPushButton('ส่งออก Excel (รายการที่กรอง)')
        filters.addWidget(reset_filters)
        filters.addWidget(export_button)
        layout.addLayout(filters)
        layout.addWidget(QLabel('คลิกขวาที่หัวคอลัมน์เพื่อเลือกค่า • กรองหลายคอลัมน์ร่วมกันได้'))
        column_filters = {}
        filter_summary = QLabel('ยังไม่ได้ใช้ตัวกรอง')
        filter_summary.setObjectName('recordsFilterSummary')
        filter_summary.setWordWrap(True)
        filter_summary.setStyleSheet('background: #e0f2fe; color: #164e63; padding: 8px; border: 1px solid #7dd3fc; font-weight: bold;')
        layout.addWidget(filter_summary)
        template = openpyxl.load_workbook(ROOT / 'templateRegistry.xlsx', read_only=True)
        mapping = {r[2]: (r[0], r[1]) for r in template['fieldmapping'].iter_rows(min_row=2, values_only=True) if len(r) >= 3 and r[2]}
        template.close()
        def column_title(key):
            return str(mapping.get(key, ('', ''))[1] or source.fields.get(key, {}).get('thai_label') or key)
        header = QTableWidget(3, len(keys) + 2)
        header.setObjectName('recordsHeader')
        header.setStyleSheet('QTableWidget { gridline-color: #829ab0; border: 1px solid #829ab0; }')
        table = QTableWidget(len(records), len(keys) + 2)
        table.setObjectName('recordsBody')
        for widget in (header, table):
            widget.setWordWrap(True)
            widget.setTextElideMode(Qt.TextElideMode.ElideNone)
            widget.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            widget.horizontalHeader().hide()
            widget.verticalHeader().hide()
            widget.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        header.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        header.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        table.horizontalScrollBar().valueChanged.connect(header.horizontalScrollBar().setValue)
        header.horizontalScrollBar().valueChanged.connect(table.horizontalScrollBar().setValue)
        widths = [144, 48] + [105 if k in ('cid', 'name', 'sch', 'hcu', 'dob', 'iv_date') else 76 for k in keys]
        for col, width in enumerate(widths):
            header.setColumnWidth(col, width)
            table.setColumnWidth(col, width)
        for col, title in enumerate(('จัดการ', 'แถว Excel')):
            header.setItem(0, col, QTableWidgetItem(title))
            header.setSpan(0, col, 3, 1)
        groups = []
        for col, key in enumerate(keys, 2):
            group, label = mapping.get(key, ('', source.fields[key]['thai_label']))
            groups.append(group)
            for row, text in enumerate((group, label, key)):
                item = QTableWidgetItem(str(text or ''))
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                item.setBackground(QColor(('#ead6db', '#e3edf6', '#d7e6f1')[row]))
                header.setItem(row, col, item)
        start = 0
        while start < len(groups):
            end = start + 1
            while end < len(groups) and groups[end] == groups[start]:
                end += 1
            if end - start > 1:
                header.setSpan(0, start + 2, 1, end - start)
            start = end
        header.resizeRowsToContents()
        for row in (1, 2):
            height = 32
            for col in range(2, header.columnCount()):
                label = QLabel(header.item(row, col).text())
                label.setWordWrap(True)
                label.setAlignment(Qt.AlignmentFlag.AlignCenter)
                label.setContentsMargins(4, 4, 4, 4)
                label.setStyleSheet('background: ' + ('#e3edf6' if row == 1 else '#d7e6f1') + '; border-right: 1px solid #829ab0; border-bottom: 1px solid #829ab0;')
                label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
                label.setText('\u200b'.join(label.text()))
                header.setCellWidget(row, col, label)
                height = max(height, label.heightForWidth(widths[col] - 10) + 12)
            header.setRowHeight(row, height)
        header.setFixedHeight(sum(header.rowHeight(r) for r in range(3)) + 4)
        layout.addWidget(header)
        layout.setSpacing(3)
        def edit_record(cid, round_id):
            if not self.discard_ok():
                return
            self.registry = Registry(source.path)
            self.dirty = False
            self.search.clear()
            self.refresh_list()
            for row, (_, record) in enumerate(self.visible_records):
                if (str(record['cid']), record['round_id']) == (cid, round_id):
                    self.table.selectRow(row)
                    self.load_selected()
                    dialog.accept()
                    self.scroll.verticalScrollBar().setValue(0)
                    return
            self.error('ไม่พบรายการ กรุณาเปิดรายการใหม่')
        def delete_record(cid, round_id, name):
            if QMessageBox.question(dialog, 'ยืนยันลบข้อมูล', f'ลบ {name}\nCID {cid} • รอบ {round_id} ?\nลบเฉพาะรอบนี้ และสำรองไฟล์ก่อนลบ', QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No) != QMessageBox.StandardButton.Yes:
                return
            if self.editing == (cid, round_id) and not self.discard_ok():
                return
            try:
                if source.path == (ROOT / 'templateRegistry.xlsx').resolve():
                    raise ValueError('ไม่ลบข้อมูลใน Template ต้นฉบับ กรุณาเลือกไฟล์ฐานข้อมูลสำเนา')
                source.delete(cid, round_id)
                self.registry = source
                remaining = source.records()
                count_label.setText(f'{len(remaining)} รอบตรวจ • {len({str(d.get("cid")) for _, d in remaining})} คน • อ่านจากไฟล์ล่าสุด')
                if self.editing == (cid, round_id):
                    self.reset_form()
                self.refresh_list()
                for row in range(table.rowCount()):
                    if table.item(row, 1).data(Qt.ItemDataRole.UserRole) == (cid, round_id):
                        table.removeRow(row)
                        break
                current = {(str(d['cid']), d['round_id']): r for r, d in source.records()}
                for row in range(table.rowCount()):
                    item = table.item(row, 1)
                    item.setText(str(current[item.data(Qt.ItemDataRole.UserRole)]))
                filter_rows()
            except Exception as exc:
                self.error(exc)
        for row, (excel_row, record) in enumerate(records):
            cid, round_id = str(record['cid']), record['round_id']
            item = QTableWidgetItem(str(excel_row))
            item.setData(Qt.ItemDataRole.UserRole, (cid, round_id))
            table.setItem(row, 1, item)
            actions = QWidget()
            buttons = QHBoxLayout(actions)
            buttons.setContentsMargins(2, 2, 2, 2)
            for label, callback in [('Edit', lambda checked=False, c=cid, r=round_id: edit_record(c, r)), ('Delete', lambda checked=False, c=cid, r=round_id, n=record.get('name'): delete_record(c, r, n))]:
                button = QPushButton(label)
                button.setStyleSheet('padding: 6px; background: ' + ('#2563a6' if label == 'Edit' else '#ba3545') + '; color: white;')
                button.clicked.connect(callback)
                buttons.addWidget(button)
            table.setCellWidget(row, 0, actions)
            for col, key in enumerate(keys, 2):
                value = record.get(key)
                if key in ('dob', 'iv_date') and isinstance(value, (date, datetime)):
                    text = value.strftime('%Y-%m-%d')
                else:
                    text = value.isoformat() if isinstance(value, (date, datetime)) else '' if value is None else str(value)
                item = QTableWidgetItem(text)
                cell = source.book['Records'].cell(excel_row, source.columns[key])
                if cell.fill.fgColor.type == 'rgb' and cell.fill.patternType == 'solid':
                    item.setBackground(QColor('#' + cell.fill.fgColor.rgb[-6:]))
                table.setItem(row, col, item)
        table.resizeRowsToContents()
        for row in range(table.rowCount()):
            table.setRowHeight(row, max(42, table.rowHeight(row)))
        def filter_rows(*args):
            visible = []
            for row in range(table.rowCount()):
                matches = any(search.text().casefold() in table.item(row, col).text().casefold() for col in range(1, table.columnCount()))
                for combo, key in ((school_filter, 'sch'), (round_filter, 'round_id')):
                    if combo.currentData() is not None:
                        matches = matches and table.item(row, keys.index(key)+2).text() == combo.currentData()
                for col, allowed in column_filters.items():
                    matches = matches and table.item(row, col).text() in allowed
                table.setRowHidden(row, not matches)
                if matches:
                    visible.append(table.item(row, 1).data(Qt.ItemDataRole.UserRole)[0])
            count_label.setText(f'แสดง {len(visible)} / {table.rowCount()} รอบตรวจ • {len(set(visible))} คน • กรองคอลัมน์ {len(column_filters)} ช่อง')
            descriptions = []
            for combo, title in ((school_filter, 'โรงเรียน'), (round_filter, 'รอบตรวจ')):
                if combo.currentData() is not None:
                    descriptions.append(f'{title}: {combo.currentText()}')
            if search.text():
                descriptions.append(f'ค้นหา: {search.text()}')
            for col, allowed in sorted(column_filters.items()):
                values = ', '.join(value or '(ว่าง)' for value in sorted(allowed)) if allowed else '(ไม่เลือกค่าใด)'
                descriptions.append(f'{column_title(keys[col-2])}: {values}')
            filter_summary.setText('กำลังกรอง • ' + ' | '.join(descriptions) if descriptions else 'ยังไม่ได้ใช้ตัวกรอง • แสดงทุกรายการ')

        def choose_column(position):
            col = header.columnAt(position.x())
            if col < 2:
                return
            chooser = QDialog(dialog)
            chooser.setWindowTitle(f'กรอง {column_title(keys[col-2])}')
            chooser.resize(360, 420)
            box = QVBoxLayout(chooser)
            title_label = QLabel(column_title(keys[col-2]))
            title_label.setWordWrap(True)
            box.addWidget(title_label)
            area = QScrollArea()
            area.setWidgetResizable(True)
            content = QWidget()
            items = QVBoxLayout(content)
            checks = []
            for value in sorted({table.item(r, col).text() for r in range(table.rowCount())}):
                check = QCheckBox(value or '(ว่าง)')
                check.setChecked(col not in column_filters or value in column_filters[col])
                checks.append((value, check))
                items.addWidget(check)
            area.setWidget(content)
            box.addWidget(area)
            for title, state in [('เลือกทั้งหมด', True), ('ไม่เลือกทั้งหมด', False)]:
                button = QPushButton(title)
                button.clicked.connect(lambda checked=False, state=state: [c.setChecked(state) for _, c in checks])
                box.addWidget(button)
            buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
            buttons.accepted.connect(chooser.accept)
            buttons.rejected.connect(chooser.reject)
            box.addWidget(buttons)
            if chooser.exec() == QDialog.DialogCode.Accepted:
                selected = {value for value, check in checks if check.isChecked()}
                if len(selected) == len(checks):
                    column_filters.pop(col, None)
                else:
                    column_filters[col] = selected
                header.cellWidget(2, col).setText(('▼ ' if col in column_filters else '') + '\u200b'.join(keys[col-2]))
                filter_rows()
        header.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        header.customContextMenuRequested.connect(choose_column)
        def clear_filters():
            column_filters.clear()
            search.clear()
            school_filter.setCurrentIndex(0)
            round_filter.setCurrentIndex(0)
            for col, key in enumerate(keys, 2):
                header.cellWidget(2, col).setText('\u200b'.join(key))
            filter_rows()
        def export_rows():
            filename, _ = QFileDialog.getSaveFileName(dialog, 'ส่งออกเฉพาะรายการที่แสดง', 'Records-export.xlsx', 'Excel (*.xlsx)')
            if not filename:
                return
            target = Path(filename).with_suffix('.xlsx')
            if target.resolve() in (source.path.resolve(), (ROOT / 'templateRegistry.xlsx').resolve()):
                self.error('กรุณาใช้ชื่อไฟล์ใหม่ ไม่ทับฐานข้อมูลหรือ Template')
                return
            try:
                from .records_export import export_records
                rows = [int(table.item(r, 1).text()) for r in range(table.rowCount()) if not table.isRowHidden(r)]
                export_records(source.book, rows, target)
                QMessageBox.information(dialog, 'ส่งออกสำเร็จ', f'ส่งออก {len(rows)} รอบตรวจแล้ว\n{target}')
            except Exception as exc:
                self.error(exc)
        export_button.clicked.connect(export_rows)
        reset_filters.clicked.connect(clear_filters)
        school_filter.currentIndexChanged.connect(filter_rows)
        round_filter.currentIndexChanged.connect(filter_rows)
        search.textChanged.connect(filter_rows)
        filter_rows()
        layout.addWidget(table)
        close = QPushButton('ปิดหน้ารายการ')
        close.clicked.connect(dialog.accept)
        layout.addWidget(close)
        dialog.exec()

    @staticmethod
    def validate_settings(values):
        if not isinstance(values, dict) or set(values) - REPEAT_FIELDS.keys():
            raise ValueError('ไฟล์ตั้งค่าข้อมูลใช้ซ้ำมีรูปแบบไม่ถูกต้อง')
        clean = {}
        for key, value in values.items():
            value = str(value).strip()
            if not value:
                raise ValueError(f'{REPEAT_FIELDS[key]}: กรุณาระบุค่าที่จะใช้ซ้ำ')
            if key == 'round_id' and (not re.fullmatch(r'[0-9]+', value) or int(value) < 1):
                raise ValueError('รอบที่ต้องเป็นจำนวนเต็มตั้งแต่ 1')
            if key == 'iv_date':
                if not re.fullmatch(r'[0-9]{2}-[0-9]{2}-[0-9]{4}', value) or int(value[-4:]) < 2400:
                    raise ValueError('วันที่ตรวจต้องเป็น DD-MM-YYYY ปี พ.ศ.')
                parse_date(value)
            clean[key] = value
        return clean

    def save_settings(self, values):
        values = self.validate_settings(values)
        path = self.registry.path.with_suffix('.settings.json')
        handle, temporary = tempfile.mkstemp(dir=path.parent, suffix='.json')
        try:
            with os.fdopen(handle, 'w', encoding='utf-8') as stream:
                json.dump({'version': 1, 'repeat_values': values}, stream, ensure_ascii=False, indent=2)
            os.replace(temporary, path)
        finally:
            Path(temporary).unlink(missing_ok=True)
        self.repeat_values = values
        for key, badge in self.repeat_badges.items():
            badge.setVisible(key in values)
            if key in values:
                badge.setToolTip('ค่าที่เติมในรายการใหม่: ' + values[key])
        self.statusBar().showMessage('บันทึกค่าใช้ซ้ำแล้ว • มีผลเมื่อกดรายการใหม่')

    def open_settings(self):
        if not self.registry:
            self.error('กรุณาเปิดฐานข้อมูลก่อนตั้งค่า')
            return
        dialog = QDialog(self)
        dialog.setWindowTitle('ตั้งค่าข้อมูลใช้ซ้ำ')
        dialog.setMinimumWidth(600)
        layout = QVBoxLayout(dialog)
        note = QLabel('เลือกฟิลด์ที่จะเติมอัตโนมัติเมื่อเริ่มรายการใหม่\nยังแก้เฉพาะรายการได้ • ไม่เปลี่ยนข้อมูลในรายการที่กำลังกรอกหรือรายการเก่า\nวันที่ตรวจที่เลือกใช้ซ้ำเป็นวันที่คงที่ กรุณาปรับเมื่อเปลี่ยนวันตรวจ')
        note.setWordWrap(True)
        layout.addWidget(note)
        form = QFormLayout()
        controls = {}
        for key, label in REPEAT_FIELDS.items():
            checkbox = QCheckBox(label)
            checkbox.setChecked(key in self.repeat_values)
            editor = QLineEdit()
            if key == 'iv_date':
                editor.setInputMask('00-00-0000;_')
            editor.setText(self.repeat_values.get(key, self.widgets[key].text()))
            editor.setEnabled(checkbox.isChecked())
            checkbox.toggled.connect(editor.setEnabled)
            form.addRow(checkbox, editor)
            controls[key] = checkbox, editor
        layout.addLayout(form)
        location = QLabel('เก็บที่: ' + str(self.registry.path.with_suffix('.settings.json')))
        location.setWordWrap(True)
        layout.addWidget(location)
        error = QLabel()
        error.setWordWrap(True)
        error.setStyleSheet('color: #b91c1c;')
        layout.addWidget(error)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel)
        def save():
            try:
                self.save_settings({k: ed.text() for k, (cb, ed) in controls.items() if cb.isChecked()})
                dialog.accept()
            except (ValueError, OSError) as exc:
                error.setText(str(exc))
        buttons.accepted.connect(save)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        dialog.exec()

    def editor(self, key):
        spec = self.registry.fields[key]
        widget = QLineEdit()
        widget.setObjectName(key)
        widget.setProperty('field_key', key)
        widget.setMinimumWidth(40)
        if spec['type'] == 'enum':
            widget.setPlaceholderText(str(spec['allowed_codes']))
        elif spec['type'] == 'date':
            widget.setInputMask('00-00-0000;_')
            widget.setToolTip('วัน-เดือน-ปี พ.ศ. เช่น 12032563 หรือ 12-03-2563')
            widget.textChanged.connect(self.update_age)
        if spec['type'] != 'date':
            widget.setToolTip(f"{spec['thai_label']}\nรหัส: {spec.get('allowed_codes') or spec['type']}")
        widget.textChanged.connect(self.changed)
        if key in ['d' + t for t in FRONT]:
            widget.textChanged.connect(self.update_plaque)
        widget.textEdited.connect(lambda text, k=key: self.input_edited(k))
        widget.returnPressed.connect(lambda k=key: self.next_input(k))
        widget.installEventFilter(self)
        self.widgets[key] = widget
        error = QLabel()
        error.setObjectName('fieldError')
        error.setWordWrap(True)
        error.hide()
        self.field_errors[key] = error
        container = QWidget()
        layout = QVBoxLayout(container)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(2)
        layout.addWidget(widget)
        layout.addWidget(error)
        return container

    @staticmethod
    def display_date(value):
        value = parse_date(value)
        return f'{value.day:02}-{value.month:02}-{value.year + 543:04}'

    def input_date(self, key):
        text = self.widgets[key].text()
        if not re.fullmatch(r'[0-9]{2}-[0-9]{2}-[0-9]{4}', text) or int(text[-4:]) < 2400:
            raise ValueError('กรอก DD-MM-YYYY โดยใช้ปี พ.ศ. เช่น 12-03-2563')
        return parse_date(text)

    def update_age(self, *args):
        if not hasattr(self, 'age_display') or not {'dob', 'iv_date'} <= self.widgets.keys():
            return
        try:
            birth, exam = self.input_date('dob'), self.input_date('iv_date')
            if exam < birth:
                raise ValueError('วันตรวจต้องไม่ก่อนวันเกิด')
            years, months = age_at(birth, exam)
            self.age_display.setText(f'{years} ปี {months} เดือน')
            if hasattr(self, 'assessment_group'):
                self.assessment_group.setText('กลุ่มอายุ: ' + ('ต่ำกว่า 3 ปี' if years < 3 else 'ตั้งแต่ 3 ปีขึ้นไป • ใช้เกณฑ์กลุ่ม 3–6 ปี'))
        except ValueError:
            self.age_display.clear()
            if hasattr(self, 'assessment_group'):
                self.assessment_group.setText('กลุ่มอายุ: รอวันที่ครบและถูกต้อง')
        for key in ('dob', 'iv_date'):
            if self.widgets[key].hasAcceptableInput():
                self.check_input(key)

    def update_plaque(self, *args):
        if self.loading_record or not hasattr(self, 'plaque_mode') or not {'vplaque', *['d' + t for t in FRONT]} <= self.widgets.keys():
            return
        if self.plaque_manual:
            self.plaque_mode.setText('ใช้ค่าที่กรอกเอง')
            return
        values = [self.widgets['d' + t].text().strip() for t in FRONT]
        if any(v not in ('0', '1', '2', '3', '9') for v in values):
            result, note = '', 'รอ Debris Index ครบ 8 ช่อง'
        elif all(v == '9' for v in values):
            result, note = '', 'ทุกซี่เป็น 9 • กรุณากรอกผลตรวจเอง'
        else:
            result = '1' if any(v in ('1', '2', '3') for v in values) else '0'
            note = 'คำนวณจาก Debris Index' + (' • ไม่รวมซี่รหัส 9' if '9' in values else '')
        self.widgets['vplaque'].setText(result)
        self.plaque_mode.setText(note)
        if result:
            self.check_input('vplaque')

    def use_auto_plaque(self):
        self.plaque_manual = False
        self.update_plaque()

    def input_edited(self, key):
        if key == 'vplaque':
            self.plaque_manual = True
            self.plaque_mode.setText('ใช้ค่าที่กรอกเอง')
        widget = self.widgets[key]
        if key.startswith('t') and key[1:].isdigit():
            position = widget.cursorPosition()
            widget.setText(widget.text().upper())
            widget.setCursorPosition(position)
        # Codes have a complete value after one keystroke; longer inputs are checked on leaving.
        if self.registry.fields[key]['type'] == 'enum' or widget.property('invalid'):
            self.check_input(key)

    def check_input(self, key):
        widget = self.widgets[key]
        spec = self.registry.fields[key]
        value = widget.text().strip()
        message = ''
        try:
            if not value:
                if key != 'dev_note' or self.widgets['dev'].text().strip() == '2' or spec.get('NotBlank') == 1:
                    raise ValueError('กรุณากรอกข้อมูล')
            elif spec['type'] == 'enum':
                if value not in str(spec['allowed_codes']).split(','):
                    raise ValueError('ใช้รหัส ' + str(spec['allowed_codes']))
            elif spec['type'] in ('integer', 'decimal'):
                number = float(value)
                if not math.isfinite(number) or number < 0 or (spec['type'] == 'integer' and not number.is_integer()):
                    raise ValueError('ต้องเป็นจำนวนเต็มไม่ติดลบ' if spec['type'] == 'integer' else 'ต้องเป็นจำนวนไม่ติดลบ')
                if key == 'round_id' and number < 1:
                    raise ValueError('รอบตรวจต้องเริ่มที่ 1')
                if key in ('wt', 'ht') and number <= (10 if key == 'wt' else 70):
                    raise ValueError('ต้องมากกว่า ' + ('10 กก.' if key == 'wt' else '70 ซม.'))
            elif spec['type'] == 'date':
                try:
                    self.input_date(key)
                except ValueError:
                    raise ValueError('วันที่ไม่ถูกต้อง ใช้ DD-MM-YYYY ปี พ.ศ. เช่น 12-03-2563')
            if key == 'cid' and value and (not re.fullmatch(r'[0-9]{13}', value) or value.startswith('00')):
                raise ValueError('ใช้เลข 13 หลัก และไม่ขึ้นต้น 00')
            if key in ('dob', 'iv_date') and value:
                other = self.widgets['dob' if key == 'iv_date' else 'iv_date'].text().strip()
                try:
                    birth = self.input_date('dob')
                    exam = self.input_date('iv_date')
                except ValueError:
                    pass
                else:
                    if other and exam < birth:
                        raise ValueError('วันตรวจต้องไม่ก่อนวันเกิด')
        except (ValueError, TypeError) as exc:
            message = str(exc)
            if message.startswith('could not convert'):
                message = 'กรุณากรอกตัวเลข'
        widget.setProperty('invalid', bool(message))
        widget.style().unpolish(widget)
        widget.style().polish(widget)
        self.field_errors[key].setText(message)
        is_tooth = key.startswith('t') and key[1:].isdigit()
        is_debris = key in ['d' + t for t in FRONT] or key == 'vplaque'
        self.field_errors[key].setVisible(bool(message) and not (is_tooth or is_debris))
        if is_tooth:
            widget.setToolTip(f"ฟัน {key[1:]}: {message or TOOTH_LABELS.get(value, '')}")
            self.refresh_tooth_errors()
        if is_debris and hasattr(self, 'debris_error'):
            errors = [f"{k[1:] if k != 'vplaque' else 'Visible plaque'}: {self.field_errors[k].text()}"
                      for k in ['vplaque'] + ['d' + t for t in FRONT] if k in self.widgets and self.widgets[k].property('invalid')]
            self.debris_error.setText(' • '.join(errors))
        return not message

    def refresh_tooth_errors(self):
        for arch, teeth in [('ฟันบน', TEETH[:10]), ('ฟันล่าง', TEETH[10:])]:
            if arch not in self.arch_errors:
                continue
            invalid = [t for t in teeth if self.widgets['t' + t].property('invalid')]
            text = ('ตรวจซี่ ' + ', '.join(invalid) + ' • กรอก A,H,N,B,C,D,E,F,G,T,9 (ฟันปกติ=A)') if invalid else ''
            self.arch_errors[arch].setText(text)
            self.arch_errors[arch].setVisible(bool(text))

    def next_input(self, key):
        if not self.check_input(key):
            return
        keys = list(self.widgets)
        index = keys.index(key)
        if index + 1 < len(keys):
            target = self.widgets[keys[index + 1]]
            target.setFocus()
            target.selectAll()
            self.scroll.ensureWidgetVisible(target, 30, 50)

    def eventFilter(self, obj, event):
        key = obj.property('field_key')
        if key in self.widgets:
            if event.type() == QEvent.Type.FocusIn:
                self.scroll.ensureWidgetVisible(obj, 30, 50)
            elif event.type() == QEvent.Type.FocusOut:
                self.check_input(key)
        return super().eventFilter(obj, event)

    def fill_debris(self):
        code = self.debris_bulk.text().strip() or '0'
        self.debris_bulk.setText(code)
        if any(code not in str(self.registry.fields['d' + t]['allowed_codes']).split(',') for t in FRONT):
            self.debris_feedback.setStyleSheet('color: #b91c1c;')
            self.debris_feedback.setText('เติมไม่ได้: ใช้รหัส 0,1,2,3,9')
            return
        count = 0
        for tooth in FRONT:
            widget = self.widgets['d' + tooth]
            if not widget.text().strip():
                widget.setText(code)
                self.check_input('d' + tooth)
                count += 1
        self.debris_feedback.setStyleSheet('color: #166534;')
        self.debris_feedback.setText(f'เติม {code} แล้ว {count} ช่อง • คงค่าซี่ที่กรอกแล้ว')

    def fill_teeth(self, arch, teeth):
        control = self.bulk_inputs[arch]
        code = control.text().strip().upper() or 'A'
        control.setText(code)
        if any(code not in str(self.registry.fields['t' + tooth]['allowed_codes']).split(',') for tooth in teeth):
            self.bulk_errors[arch].setStyleSheet('color: #b91c1c;')
            self.bulk_errors[arch].setText('เติมไม่ได้: ใช้ A สำหรับฟันปกติ หรือ H,N,B,C,D,E,F,G,T,9 • ไม่ใช้ 0')
            self.bulk_errors[arch].show()
            control.setFocus()
            return
        self.bulk_errors[arch].setText('')
        count = 0
        for tooth in teeth:
            widget = self.widgets['t' + tooth]
            if not widget.text().strip():
                widget.setText(code)
                self.check_input('t' + tooth)
                count += 1
        self.bulk_errors[arch].show()
        self.bulk_errors[arch].setStyleSheet('color: #166534;')
        self.bulk_errors[arch].setText(f'เติม {code} แล้ว {count} ช่อง • คงค่าซี่ที่กรอกแล้ว' if count else 'ไม่มีช่องว่าง • หากต้องการแก้รหัส ให้แก้ที่ซี่ฟันโดยตรง')
        self.statusBar().showMessage(f'เติม{arch} {count} ช่อง • คงค่าซี่ที่กรอกแล้ว')

    def add_section(self, layout, title, keys):
        box = QGroupBox(title)
        grid = QGridLayout(box)
        grid.setSpacing(0)
        grid.setColumnStretch(1, 3)
        grid.setColumnStretch(2, 2)
        for col, text in enumerate(('field', 'หัวข้อคำถาม', 'ค่าที่กรอก')):
            label = QLabel(text)
            label.setObjectName('sheetHeader')
            grid.addWidget(label, 0, col)
        row = 1
        for key in keys:
            field = QLabel(key)
            field.setObjectName('fieldName')
            label = QLabel({'dob': 'วันเดือนปีเกิด (DD-MM-YYYY พ.ศ.)', 'iv_date': 'วันที่ตรวจ (DD-MM-YYYY พ.ศ.)'}.get(key, self.registry.fields[key]['thai_label']))
            label.setObjectName('sheetLabel')
            label.setWordWrap(True)
            grid.addWidget(field, row, 0)
            if key in REPEAT_FIELDS:
                labels = QWidget()
                labels_layout = QVBoxLayout(labels)
                labels_layout.setContentsMargins(0, 0, 0, 0)
                labels_layout.setSpacing(0)
                labels_layout.addWidget(label)
                badge = QLabel('ใช้ซ้ำ • เติมเมื่อเริ่มรายการใหม่')
                badge.setStyleSheet('color: #176b59; background: #e5f5ef; padding: 2px 6px; font-size: 10px;')
                badge.setVisible(key in self.repeat_values)
                labels_layout.addWidget(badge)
                self.repeat_badges[key] = badge
                grid.addWidget(labels, row, 1)
            else:
                grid.addWidget(label, row, 1)
            grid.addWidget(self.editor(key), row, 2)
            row += 1
            if key == 'iv_date':
                self.age_display = QLineEdit()
                self.age_display.setReadOnly(True)
                self.age_display.setFocusPolicy(Qt.FocusPolicy.NoFocus)
                self.age_display.setPlaceholderText('คำนวณเมื่อกรอกวันที่ครบ')
                grid.addWidget(QLabel('age_y'), row, 0)
                grid.addWidget(QLabel('อายุ (ปี) / เศษเดือน'), row, 1)
                grid.addWidget(self.age_display, row, 2)
                row += 1
        layout.addWidget(box)

    def build_form(self):
        old = self.scroll.takeWidget()
        if old:
            old.deleteLater()
        self.widgets = {}
        self.field_errors = {}
        self.repeat_badges = {}
        self.bulk_inputs = {}
        self.bulk_errors = {}
        self.bulk_buttons = {}
        self.arch_errors = {}
        page = QWidget()
        page.setMinimumWidth(1000)
        page.setObjectName("entrySheet")
        columns = QHBoxLayout(page)
        columns.setSpacing(18)
        left, right = QVBoxLayout(), QVBoxLayout()
        right.setSpacing(3)
        columns.addLayout(left, 1)
        columns.addLayout(right, 1)
        keys = list(self.registry.fields)
        self.add_section(left, '1 ข้อมูลทั่วไป', [k for k in keys[:21] if k not in ('age_y', 'age_m')])
        self.add_section(left, '2 บริการส่งเสริมป้องกัน', keys[21:24])
        self.add_section(left, '3 พฤติกรรมดูแลช่องปาก', keys[24:40])
        self.add_section(left, '4 ปัญหาช่องปาก', keys[40:46])
        left.addStretch()
        hint = QLabel('พิมพ์รหัสแล้วกด Enter หรือ Tab ไปช่องถัดไป\nค่าที่ผิดจะแสดงกรอบแดงและคำอธิบายใต้ช่อง')
        hint.setWordWrap(True)
        right.addWidget(hint)
        legend = QGroupBox('ความหมายรหัสฟัน')
        legend_layout = QGridLayout(legend)
        legend_layout.setColumnStretch(1, 1)
        for row, (code, label) in enumerate(TOOTH_LABELS.items()):
            legend_layout.addWidget(QLabel(code), row, 0)
            legend_layout.addWidget(QLabel(label), row, 1)
        right.addWidget(legend)
        for title, teeth in [('ฟันบน', TEETH[:10]), ('ฟันล่าง', TEETH[10:])]:
            box = QGroupBox(title)
            grid = QGridLayout(box)
            grid.setSpacing(0)
            for col, tooth in enumerate(teeth):
                label = QLabel(tooth)
                label.setObjectName("toothNumber")
                label.setAlignment(Qt.AlignmentFlag.AlignCenter)
                grid.addWidget(label, 0, col)
                grid.addWidget(self.editor('t' + tooth), 1, col)
                widget = self.widgets['t' + tooth]
                widget.setPlaceholderText('')
                widget.setAlignment(Qt.AlignmentFlag.AlignCenter)
                widget.setProperty('toothCell', True)
                widget.setFixedHeight(32)
            controls = QWidget()
            controls_layout = QHBoxLayout(controls)
            controls_layout.setContentsMargins(10, 0, 0, 0)
            code = QLineEdit()
            code.setPlaceholderText('A')
            code.setToolTip('ฟันปกติใช้ A; รหัสฟันไม่ใช้ 0')
            code.setFixedWidth(48)
            code.setAlignment(Qt.AlignmentFlag.AlignCenter)
            button = QPushButton('เติมช่องว่าง')
            button.setObjectName('fillButton')
            button.setToolTip('เติมเฉพาะช่องว่างของ' + title)
            button.clicked.connect(lambda checked=False, arch=title, ts=teeth: self.fill_teeth(arch, ts))
            code.returnPressed.connect(lambda arch=title, ts=teeth: self.fill_teeth(arch, ts))
            controls_layout.addWidget(code)
            controls_layout.addWidget(button)
            grid.addWidget(controls, 1, 10)
            warning = QLabel()
            warning.setObjectName('fieldError')
            warning.setWordWrap(True)
            warning.setFixedHeight(24)
            warning.hide()
            grid.addWidget(warning, 2, 0, 1, 11)
            summary = QLabel()
            summary.setObjectName('fieldError')
            summary.setWordWrap(True)
            summary.setFixedHeight(24)
            summary.hide()
            grid.addWidget(summary, 3, 0, 1, 11)
            self.arch_errors[title] = summary
            self.bulk_inputs[title] = code
            self.bulk_errors[title] = warning
            self.bulk_buttons[title] = button
            right.addWidget(box)
        box = QGroupBox('คราบจุลินทรีย์')
        grid = QGridLayout(box)
        visible = QWidget()
        visible_layout = QVBoxLayout(visible)
        visible_layout.addWidget(QLabel('อายุ ≥3 ปี • Visible Plaque'))
        visible_layout.addWidget(self.editor('vplaque'))
        self.widgets['vplaque'].setFixedSize(58, 32)
        visible_layout.addWidget(QLabel('รหัส 0,1'))
        self.plaque_mode = QLabel('รอ Debris Index ครบ 8 ช่อง')
        self.plaque_mode.setWordWrap(True)
        self.plaque_mode.setMaximumWidth(160)
        visible_layout.addWidget(self.plaque_mode)
        self.plaque_auto = QPushButton('ใช้ค่าอัตโนมัติ')
        self.plaque_auto.clicked.connect(self.use_auto_plaque)
        visible_layout.addWidget(self.plaque_auto)
        grid.addWidget(visible, 0, 0, alignment=Qt.AlignmentFlag.AlignTop)
        panel = QWidget()
        cells = QGridLayout(panel)
        cells.setContentsMargins(0, 0, 0, 0)
        cells.setSpacing(0)
        heading = QLabel('อายุ <3 ปี • Simplified Debris Index')
        cells.addWidget(heading, 0, 0, 1, 4)
        for i, tooth in enumerate(FRONT):
            row, col = divmod(i, 4)
            label = QLabel(tooth)
            label.setObjectName('toothNumber')
            label.setAlignment(Qt.AlignmentFlag.AlignCenter)
            cells.addWidget(label, 1 if row == 0 else 4, col)
            cells.addWidget(self.editor('d' + tooth), 2 + row, col)
            editor = self.widgets['d' + tooth]
            editor.setFixedSize(52, 32)
            editor.setPlaceholderText('')
            editor.setProperty('toothCell', True)
            editor.setAlignment(Qt.AlignmentFlag.AlignCenter)
        cells.addWidget(QLabel('รหัส 0,1,2,3,9'), 5, 0, 1, 4)
        panel.setFixedWidth(208)
        grid.addWidget(panel, 0, 1, alignment=Qt.AlignmentFlag.AlignTop)
        controls = QWidget()
        controls_layout = QHBoxLayout(controls)
        controls_layout.setContentsMargins(6, 0, 0, 0)
        self.debris_bulk = QLineEdit('0')
        self.debris_bulk.setFixedWidth(42)
        self.debris_bulk.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.debris_bulk.setToolTip('รหัส Debris Index: 0,1,2,3,9 • ค่าเริ่มต้น 0')
        self.debris_fill_button = QPushButton('เติมช่องว่าง')
        self.debris_fill_button.setObjectName('fillButton')
        self.debris_fill_button.clicked.connect(self.fill_debris)
        self.debris_bulk.returnPressed.connect(self.fill_debris)
        controls_layout.addWidget(self.debris_bulk)
        controls_layout.addWidget(self.debris_fill_button)
        grid.addWidget(controls, 0, 2, alignment=Qt.AlignmentFlag.AlignVCenter)
        self.debris_error = QLabel()
        self.debris_error.setObjectName('fieldError')
        self.debris_error.setWordWrap(True)
        self.debris_error.setMinimumHeight(38)
        grid.addWidget(self.debris_error, 1, 0, 1, 3)
        self.debris_feedback = QLabel()
        self.debris_feedback.setMinimumHeight(22)
        grid.addWidget(self.debris_feedback, 2, 0, 1, 3)
        legend = QLabel(
            'ความหมายรหัส Debris Index\n'
            '0  ไม่มีคราบเศษอาหาร\n'
            '1  คราบอ่อนปกคลุมผิวฟันน้อยกว่า 1/3\n'
            '2  คราบอ่อนปกคลุมผิวฟัน 1/3–2/3\n'
            '3  คราบอ่อนปกคลุมผิวฟันมากกว่า 2/3\n'
            '9  ไม่ได้บันทึก')
        legend.setWordWrap(True)
        legend.setStyleSheet('background: #edf3f8; color: #244b68; padding: 8px;')
        grid.addWidget(legend, 3, 0, 1, 3)
        right.addWidget(box)
        self.details = QTableWidget(0, 2)
        self.details.setHorizontalHeaderLabels(['ผลประเมิน / รายการคำนวณ', 'ผล'])
        self.details.horizontalHeader().setStretchLastSection(True)
        self.details.setColumnWidth(0, 350)
        self.details.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.details.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.details.setMinimumHeight(620)
        output_box = QGroupBox('ผลประเมินอัตโนมัติ • อ่านอย่างเดียว')
        output_box.setObjectName('assessmentBox')
        output_layout = QVBoxLayout(output_box)
        self.assessment_group = QLabel('กลุ่มอายุ: รอวันเกิดและวันที่ตรวจ')
        self.assessment_group.setWordWrap(True)
        self.assessment_group.setStyleSheet('color: #6b3c9b; font-weight: bold; padding: 8px;')
        output_layout.addWidget(self.assessment_group)
        self.details.setObjectName('assessmentTable')
        output_layout.addWidget(self.details)
        right.addWidget(output_box)
        right.addStretch()
        self.scroll.setWidget(page)
        ordered = list(self.widgets.values())
        for first, second in zip(ordered, ordered[1:]):
            QWidget.setTabOrder(first, second)

    def reset_form(self):
        self.loading_record = True
        self.plaque_manual = False
        self.editing = None
        self.assessment_timer.stop()
        self.debris_error.clear()
        self.debris_bulk.setText('0')
        self.debris_feedback.clear()
        for arch, control in self.bulk_inputs.items():
            control.setText('A')
            self.bulk_errors[arch].setText('')
            self.arch_errors[arch].setText('')
            self.arch_errors[arch].hide()
            self.bulk_errors[arch].hide()
        for key, widget in self.widgets.items():
            widget.clear()
            widget.setProperty('invalid', False)
            widget.style().unpolish(widget)
            widget.style().polish(widget)
            self.field_errors[key].hide()
        if self.widgets:
            self.widgets['round_id'].setText('1')
            today = date.today()
            self.widgets['iv_date'].setText(f'{today.day:02}-{today.month:02}-{today.year + 543}')
        for key, value in self.repeat_values.items():
            self.widgets[key].setText(value)
        self.update_age()
        self.loading_record = False
        self.update_plaque()
        self.mode.setText('รายการใหม่ • ต้องกรอกทุกช่อง ยกเว้นหมายเหตุเมื่อพัฒนาการปกติ')
        self.result.setText('เกณฑ์: <3 ปี / ≥3 ปี (รวมอายุ 6 ปีขึ้นไปตามที่กำหนด)')
        self.result.setStyleSheet('')
        if hasattr(self, 'details'):
            self.details.setRowCount(0)
        self.dirty = False

    def new_record(self):
        if self.registry and self.discard_ok():
            self.reset_form()
            self.scroll.verticalScrollBar().setValue(0)
            self.widgets['round_id'].setFocus()

    def refresh_list(self):
        if not self.registry:
            return
        query = self.search.text().strip().casefold()
        records = [(r, d) for r, d in self.registry.records() if query in ' '.join(str(d.get(k, '')) for k in ('cid', 'name', 'sch')).casefold()]
        self.visible_records = records
        self.table.setRowCount(len(records))
        for i, (_, d) in enumerate(records):
            for j, key in enumerate(('name', 'round_id', 'riskscore', 'risklevel')):
                value = d.get(key)
                item = QTableWidgetItem('' if value is None else str(value))
                item.setBackground(QColor('#' + COLORS.get(d.get('risklevel'), 'FFFFFF')))
                item.setToolTip(f"CID: {d['cid']}\n{d.get('sch', '')}")
                self.table.setItem(i, j, item)
        self.count.setText(f'แสดง {len(records)} รอบตรวจ • ดับเบิลคลิกเพื่อแก้ไข')

    def load_selected(self, *args):
        index = self.table.currentRow()
        if index < 0 or not self.discard_ok():
            return
        _, d = self.visible_records[index]
        self.loading_record = True
        for key, widget in self.widgets.items():
            value = d.get(key)
            if key in ('dob', 'iv_date') and value:
                value = self.display_date(value)
            widget.setText('' if value is None else str(value))
            self.check_input(key)
        self.update_age()
        self.loading_record = False
        self.plaque_manual = True
        self.plaque_mode.setText('ใช้ค่าที่บันทึกไว้ • แก้เองได้')
        self.editing = (str(d['cid']), d['round_id'])
        self.mode.setText(f"แก้ไข: {d.get('name')} • CID {d['cid']} • รอบ {d['round_id']}")
        self.show_result(d)
        self.dirty = False

    def calculate(self):
        if not self.registry:
            raise ValueError('กรุณาตั้งค่าแหล่งข้อมูลก่อน')
        invalid = [key for key in self.widgets if not self.check_input(key)]
        if invalid:
            first = self.widgets[invalid[0]]
            first.setFocus()
            self.scroll.ensureWidgetVisible(first, 30, 50)
            raise ValueError(f'กรุณาแก้ไข {len(invalid)} ช่องที่มีกรอบแดงก่อนบันทึก')
        raw = {key: widget.text().strip().upper() if key.startswith('t') and key[1:].isdigit() else widget.text() for key, widget in self.widgets.items()}
        data = validate_record(raw, self.registry.fields)
        return assess(data, self.registry.book)

    def show_result(self, data):
        self.result.setText(f"อายุ {data.get('age_y')} ปี {data.get('age_m')} เดือน  |  {data.get('riskscore')} คะแนน  |  {data.get('risklevel')}")
        if any(str(data.get('t' + tooth)) == '9' for tooth in TEETH):
            self.result.setText(self.result.text() + '\nมีรหัสฟัน 9: จำนวนฟันระบุไม่ได้ และผลคะแนนอาจไม่ครอบคลุมซี่ที่ไม่ได้บันทึก')
        self.result.setStyleSheet('background: #' + COLORS.get(data.get('risklevel'), 'EEEEEE') + '; padding: 14px; border-radius: 8px; color: #172033;')
        calculated = [key for key in self.registry.fields if key.startswith(('calc_', 'r_'))]
        self.details.setRowCount(len(calculated))
        active_col = 3 if data.get('age_y', 0) < 3 else 4
        score_sheet = self.registry.book['ตารางRiskScore']
        active = {score_sheet.cell(r, 1).value for r in range(5, 16) if score_sheet.cell(r, active_col).value == '/'}
        for row, key in enumerate(calculated):
            label = self.registry.fields[key]['thai_label']
            if key == 'r_ohplaque':
                label = 'พบคราบจุลินทรีย์ Front-8 (พบ=3, ไม่พบ=0)'
            count_labels = {
                'calc_teeth_in_mouth': 'จำนวนซี่ฟันในช่องปาก',
                'calc_white_discolor': 'จำนวนฟันขาวขุ่น/เปลี่ยนสี (H)',
                'calc_caries_hole': 'จำนวนฟันผุเป็นรู (B, C)',
                'calc_filled': 'จำนวนฟันอุด (C, D)',
                'calc_extracted': 'จำนวนฟันหายเพราะผุ (E)',
                'calc_caries_front8': 'จำนวนฟันหน้า 8 ซี่ที่ผุเป็นรู (B, C)',
            }
            label = count_labels.get(key, label)
            self.details.setItem(row, 0, QTableWidgetItem(label))
            value = data.get(key)
            text = 'ระบุไม่ได้' if value is None else str(value)
            if key.startswith('r_') and key not in active:
                text += ' (ไม่นับในกลุ่มอายุนี้)'
            self.details.setItem(row, 1, QTableWidgetItem(text))
        self.details.resizeRowsToContents()
        self.details.setFixedHeight(self.details.horizontalHeader().height() + sum(self.details.rowHeight(r) for r in range(self.details.rowCount())) + 24)

    def preview(self):
        try:
            data = self.calculate()
            self.show_result(data)
            self.scroll.ensureWidgetVisible(self.details, 20, 20)
        except Exception as exc:
            self.error(exc)

    def save_record(self):
        try:
            data = self.calculate()
            key = (data['cid'], data['round_id'])
            if self.editing and key != self.editing:
                raise ValueError('ขณะแก้ไขห้ามเปลี่ยน CID หรือรอบตรวจ ใช้ “รายการใหม่” เพื่อบันทึกรอบใหม่')
            if self.registry.path == (ROOT / 'templateRegistry.xlsx').resolve():
                raise ValueError('ไฟล์นี้เป็น Template ต้นฉบับ กรุณาทำสำเนาไฟล์ด้วยชื่อใหม่ แล้วเลือกใน “ตั้งค่าแหล่งข้อมูล”')
            row, backup = self.registry.save(data, allow_update=self.editing == key)
            self.refresh_list()
            self.reset_form()
            self.mode.setText(f"✓ บันทึกเรียบร้อยแล้ว: {data['name']} • รอบ {data['round_id']} • แถว {row}\nพร้อมกรอกข้อมูลเด็กคนถัดไป")
            self.result.setText(f"บันทึกล่าสุด: {data['riskscore']} คะแนน • {data['risklevel']} | เริ่มรายการใหม่แล้ว")
            self.scroll.verticalScrollBar().setValue(0)
            target = self.widgets['name']
            target.setFocus()
            self.statusBar().showMessage(f'บันทึกเรียบร้อยแล้ว • สำรอง: {backup.name}')
        except Exception as exc:
            self.error(exc)


def configure_app(app):
    class ButtonStyleFilter(QWidget):
        def eventFilter(self, obj, event):
            if event.type() == QEvent.Type.Show and isinstance(obj, QPushButton):
                obj.setCursor(Qt.CursorShape.PointingHandCursor)
                if obj.text() in ('ประเมิน', 'ตั้งค่าข้อมูลใช้ซ้ำ'):
                    obj.setProperty('actionColor', 'purple')
                elif obj.text() in ('รายการใหม่', 'ดูเด็กทั้งหมด (Records)'):
                    obj.setProperty('actionColor', 'orange')
                obj.style().unpolish(obj)
                obj.style().polish(obj)
            return False
    app.button_style_filter = ButtonStyleFilter()
    app.installEventFilter(app.button_style_filter)
    # Explicit registration also supports Qt offscreen visual verification.
    font_path = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts' / 'tahoma.ttf'
    if font_path.exists():
        QFontDatabase.addApplicationFont(str(font_path))
    app.setFont(QFont('Tahoma', 10))
    app.setStyleSheet('''
        QMainWindow { background: #eef3f8; }
        QWidget#entrySheet { background: #f8fafc; }
        QLabel#sheetHeader { background: #dceaf5; color: #224762; font-weight: bold; padding: 6px; border: 1px solid #b8cad8; }
        QLabel#fieldName { background: #edf3f8; color: #52677b; font-size: 10px; padding: 4px; border-bottom: 1px solid #ced9e3; }
        QLabel#sheetLabel { background: #ffffff; padding: 4px 6px; border-bottom: 1px solid #ced9e3; }
        QLabel#toothNumber { color: #387bb7; font-weight: bold; padding: 5px 0; }
        QLineEdit[toothCell="true"] { border-radius: 0; border: 1px solid #95aabb; padding: 5px 1px; font-size: 14px; }
        QPushButton#fillButton { background: #d2eaf5; color: #20577e; border: 1px solid #7ba6c3; padding: 7px; }
        QHeaderView::section { background: #dceaf5; padding: 5px; border: 1px solid #b8cad8; color: #224762; }
        QLabel#title { font-size: 26px; font-weight: bold; color: #123c4a; }
        QLabel#mode { color: #14532d; background: #dcfce7; border: 2px solid #86c99c; border-left: 8px solid #228447; border-radius: 8px; padding: 12px 16px; font-size: 18px; font-weight: bold; }
        QLineEdit { background: white; border: 1px solid #cbd5e1; border-radius: 2px; padding: 5px; color: #172033; }
        QGroupBox#assessmentBox { background: #f1eafa; border: 2px solid #9b7aba; }
        QTableWidget#assessmentTable { background: #f8f4fc; color: #553477; }
        QTableWidget#assessmentTable QHeaderView::section { background: #ded0ed; color: #553477; }
        QPushButton { background: #2563a6; color: white; border: 0; border-radius: 6px; padding: 10px 16px; }
        QPushButton:hover { background: #347abd; }
        QPushButton:pressed { padding-top: 13px; padding-bottom: 7px; border: 2px inset #183d60; }
        QPushButton:disabled { background: #d5dbe1; color: #78828c; }
        QPushButton[actionColor="purple"] { background: #7853a5; }
        QPushButton[actionColor="orange"] { background: #b6651d; }
        QPushButton#primary { background: #087f8c; color: white; font-weight: bold; }
        QLineEdit[invalid="true"] { border: 2px solid #dc2626; background: #fff1f2; }
        QLabel#fieldError { color: #b91c1c; font-size: 11px; }
        QGroupBox { background: #ffffff; border: 1px solid #c5d4e0; border-radius: 4px; margin-top: 18px; padding: 14px 6px 6px; }
        QGroupBox::title { subcontrol-origin: margin; subcontrol-position: top left; padding: 3px 9px; color: #244b68; background: #e0ecf6; font-weight: bold; }
        QTableWidget { background: white; color: #172033; gridline-color: #e2e8f0; }
    ''')


def main():
    app = QApplication(sys.argv)
    configure_app(app)
    if len(sys.argv)>2 and sys.argv[1]=='--smoke-test':
        def smoke_warning(parent, title, message, *args, **kwargs):
            raise RuntimeError(f'{title}: {message}')
        QMessageBox.warning = smoke_warning
        resource=ROOT/'templateRegistry.xlsx'
        registry=Registry(resource)
        assert registry.records()==[]
        window=Window(resource)
        window.show();app.processEvents()
        destination=Path(sys.argv[2])
        window.grab().save(str(destination.with_suffix('.png')))
        destination.write_text(json.dumps({'ok':True,'fields':len(registry.fields),'sheets':registry.book.sheetnames},ensure_ascii=False),encoding='utf-8')
        window.dirty=False;window.close()
        return 0
    path = sys.argv[1] if len(sys.argv) > 1 else QSettings('StudentOralRegistry', 'Desktop').value('data_source', None)
    if not path:
        from .runtime import initial_database
        try:
            path=str(initial_database())
            QSettings('StudentOralRegistry', 'Desktop').setValue('data_source',path)
        except OSError as exc:
            QMessageBox.warning(None,'เริ่มต้นฐานข้อมูล',str(exc))
    window = Window(path)
    window.show()
    return app.exec()


if __name__ == '__main__':
    sys.exit(main())
