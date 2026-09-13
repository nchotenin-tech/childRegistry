from openpyxl import load_workbook
from oral_registry.records_export import export_records


def test_export_subset(cohort, tmp_path):
    destination = tmp_path / 'export.xlsx'
    export_records(cohort.book, [5, 7], destination)
    result = load_workbook(destination)
    sheet = result['Records']
    assert sheet.max_row == 5
    assert sheet.cell(4, cohort.columns['cid']).value == cohort.book['Records'].cell(5, cohort.columns['cid']).value
    assert sheet.cell(5, cohort.columns['round_id']).value == cohort.book['Records'].cell(7, cohort.columns['round_id']).value
    assert sheet.freeze_panes == 'A4'
    assert sheet.cell(3, 2).alignment.wrap_text
    assert len(cohort.records()) == 9


def test_combined_filters(cohort, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtCore import QPoint
    from PySide6.QtWidgets import QApplication, QDialog, QTableWidget, QComboBox, QCheckBox, QPushButton
    from oral_registry.app import Window, configure_app
    app = QApplication.instance() or QApplication([])
    configure_app(app)
    window = Window(cohort.path)
    def inspect(dialog):
        if dialog.windowTitle().startswith('กรอง '):
            for check in dialog.findChildren(QCheckBox):
                check.setChecked(check.text() == '1')
            return QDialog.DialogCode.Accepted
        table = dialog.findChild(QTableWidget, 'recordsBody')
        header = dialog.findChild(QTableWidget, 'recordsHeader')
        combo = dialog.findChild(QComboBox, 'filter_round_id')
        combo.setCurrentIndex(combo.findData('2'))
        visible = lambda: sum(not table.isRowHidden(r) for r in range(table.rowCount()))
        assert visible() == 3
        col = list(cohort.columns).index('sweet_freq') + 2
        header.customContextMenuRequested.emit(QPoint(header.columnViewportPosition(col)+2, 10))
        assert visible() == 1
        col = list(cohort.columns).index('round_id') + 2
        header.customContextMenuRequested.emit(QPoint(header.columnViewportPosition(col)+2, 10))
        assert visible() == 0
        next(b for b in dialog.findChildren(QPushButton) if b.text() == 'ล้างตัวกรอง').click()
        assert visible() == 9
        dialog.show()
        app.processEvents()
        dialog.grab().save('tmp/records-filters.png')
        dialog.close()
        return 0
    monkeypatch.setattr(QDialog, 'exec', inspect)
    try:
        window.show_records()
    finally:
        window.dirty = False
        window.close()
