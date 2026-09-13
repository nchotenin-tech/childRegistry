from datetime import date
from pathlib import Path
import shutil

import openpyxl
import pytest

from oral_registry.core import FRONT, TEETH, age_at, alert_matches, assess, parse_date, validate_record
from oral_registry.storage import Registry

TEMPLATE = Path(__file__).resolve().parents[1] / 'templateRegistry.xlsx'


def sample(registry, years=4):
    raw = {}
    for key, spec in registry.fields.items():
        kind = spec['type']
        raw[key] = str(spec['allowed_codes']).split(',')[0] if kind == 'enum' else '1' if kind in ('integer', 'decimal') else 'ทดสอบ'
    raw.update(cid='1234567890128', name='เด็กทดสอบ', round_id=1, dob=f'{2026-years}-09-11', iv_date='2026-09-11', wt=16.7, ht=101,
               br_method=2, br_freq=2, tm_morn=2, tm_bed=2, ftp_freq=2, bottle=4, dev=1)
    for tooth in TEETH:
        raw['t' + tooth] = 'A'
    for tooth in FRONT:
        raw['d' + tooth] = 0
    return validate_record(raw, registry.fields)


def test_dates_and_age():
    assert parse_date('11-09-2569') == date(2026, 9, 11)
    assert parse_date('2026-09-11') == date(2026, 9, 11)
    assert age_at(date(2023, 9, 12), date(2026, 9, 11)) == (2, 11)
    with pytest.raises(ValueError):
        parse_date('31-02-2569')


@pytest.mark.parametrize('score,level', [(0,'เสี่ยงต่ำ'),(2,'เสี่ยงต่ำ'),(3,'เสี่ยงสูง'),(5,'เสี่ยงสูง'),(6,'เสี่ยงสูงมาก')])
def test_thresholds(registry, score, level):
    data = sample(registry)
    registry.book['ตารางRiskScore']['G17'] = score
    result = assess(data, registry.book)
    assert result['riskscore'] == score
    assert result['risklevel'] == level


def test_age_groups_and_plaque(registry):
    data = sample(registry, 2)
    data['d52'] = 1
    assert assess(data, registry.book)['riskscore'] == 3
    data['age_y'] = 3
    assert assess(data, registry.book)['riskscore'] == 0
    data['vplaque'] = 1
    assert assess(data, registry.book)['riskscore'] == 3
    data['age_y'] = 8
    assert assess(data, registry.book)['riskscore'] == 3
    registry.book['ตารางRiskScore']['D6'] = '/'
    assert assess(data, registry.book)['riskscore'] == 6


def test_tooth_codes(registry):
    data = sample(registry)
    data.update(t55='H', t54='N', t53='B', t52='C', t51='D', t61='E')
    result = assess(data, registry.book)
    assert result['calc_white_discolor'] == 1
    assert result['calc_caries_hole'] == 2
    assert result['calc_filled'] == 2
    assert result['calc_extracted'] == 1
    assert result['calc_teeth_in_mouth'] == 19
    assert result['r_caries'] == 3
    data['t62'] = '9'
    assert assess(data, registry.book)['calc_teeth_in_mouth'] is None


@pytest.mark.parametrize('field,value', [('cid','0012345678901'),('cid','123'),('round_id',0),('wt',10),('ht',70),('sex',9),('dob','2027-01-01'),('br_freq',''),('sdrink_perday',1.5)])
def test_invalid_data(registry, field, value):
    data = sample(registry)
    data[field] = value
    with pytest.raises(ValueError):
        validate_record(data, registry.fields)


@pytest.mark.parametrize('value,rule,expected', [(0,0,True),(2,'1,2',True),(3,'>2',True),(2,'>2',False),('เสี่ยงสูง','"เสี่ยงสูง","เสี่ยงสูงมาก"',True),(None,0,False)])
def test_alerts(value, rule, expected):
    assert alert_matches(value, rule) == expected


def test_save_update_preserve_and_conflict(registry):
    data = sample(registry)
    data['name'] = '=1+1'
    result = assess(data, registry.book)
    original_count = len(registry.records())
    row, backup = registry.save(result)
    assert row >= 4 and backup.exists()
    assert len(registry.records()) == original_count + 1
    assert registry.book['Records'].cell(row, registry.columns['name']).data_type == 's'
    assert registry.book['Records'].cell(row, registry.columns['cid']).number_format == '@'
    with pytest.raises(ValueError, match='มี CID'):
        registry.save(result)
    result['br_freq'] = 0
    result = assess(result, registry.book)
    registry.save(result, allow_update=True)
    assert len(registry.records()) == original_count + 1
    cell = registry.book['Records'].cell(row, registry.columns['br_freq'])
    assert cell.fill.fgColor.rgb.endswith('FF9999')
    original = openpyxl.load_workbook(TEMPLATE)
    for name in ('Entry', 'Codebook', 'ตารางRiskScore', 'Config'):
        assert list(original[name].values) == list(registry.book[name].values)
    registry.path.write_bytes(registry.path.read_bytes() + b' ')
    with pytest.raises(ValueError, match='ภายนอก'):
        registry.save(result, allow_update=True)


def test_excel_lock(registry):
    registry.path.with_name('~$' + registry.path.name).touch()
    with pytest.raises(ValueError, match='Excel'):
        registry.save(assess(sample(registry), registry.book))


def test_ui_fields(registry, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    from oral_registry.app import Window, configure_app
    app = QApplication.instance() or QApplication([])
    configure_app(app)
    window = Window(registry.path)
    expected = {key for key in registry.fields if not key.startswith(('calc_', 'r_')) and key not in ('age_y','age_m','riskscore','risklevel')}
    assert set(window.widgets) == expected
    assert window.scroll.widget() is not None
    window.show()
    app.processEvents()
    out = Path('tmp')
    out.mkdir(exist_ok=True)
    window.grab().save(str(out / 'app-preview.png'))
    from PySide6.QtWidgets import QLineEdit
    window.widgets['t55'].setText('B')
    window.bulk_inputs['ฟันบน'].setText('a')
    window.bulk_buttons['ฟันบน'].click()
    assert window.widgets['t55'].text() == 'B'
    assert all(window.widgets['t' + t].text() == 'A' for t in TEETH[1:10])
    assert all(not window.widgets['t' + t].text() for t in TEETH[10:])
    window.bulk_inputs['ฟันล่าง'].setText('X')
    window.bulk_buttons['ฟันล่าง'].click()
    assert not window.bulk_errors['ฟันล่าง'].isHidden()
    assert all(not window.widgets['t' + t].text() for t in TEETH[10:])
    window.bulk_inputs['ฟันล่าง'].setText('9')
    window.bulk_buttons['ฟันล่าง'].click()
    assert all(window.widgets['t' + t].text() == '9' for t in TEETH[10:])
    window.reset_form()
    data = sample(registry)
    for key, widget in window.widgets.items():
        value = data.get(key)
        assert isinstance(widget, QLineEdit)
        widget.setText(window.display_date(value) if key in ('dob', 'iv_date') else '' if value is None else str(value))
    from PySide6.QtTest import QTest
    from PySide6.QtCore import Qt
    window.widgets['t55'].setFocus()
    window.widgets['t55'].selectAll()
    QTest.keyClicks(window.widgets['t55'], 'x')
    assert window.widgets['t55'].property('invalid')
    assert window.field_errors['t55'].text()
    QTest.keyClick(window.widgets['t55'], Qt.Key.Key_Return)
    assert window.widgets['t55'].hasFocus()
    with pytest.raises(ValueError, match='กรอบแดง'):
        window.calculate()
    window.widgets['t55'].selectAll()
    QTest.keyClicks(window.widgets['t55'], 'a')
    assert window.widgets['t55'].text() == 'A'
    assert not window.widgets['t55'].property('invalid')
    QTest.keyClick(window.widgets['t55'], Qt.Key.Key_Return)
    assert window.widgets['t54'].hasFocus()
    window.widgets['wt'].setText('9')
    assert not window.check_input('wt')
    window.widgets['wt'].setText('16.7')
    assert window.check_input('wt')
    assert window.calculate()['riskscore'] == 0
    window.auto_assess()
    assert '0 คะแนน' in window.result.text()
    window.save_record()
    assert window.editing is None
    assert window.widgets['name'].text() == ''
    assert 'บันทึกเรียบร้อยแล้ว' in window.mode.text()
    assert any(str(d['cid']) == data['cid'] for _, d in window.registry.records())
    assert not window.dirty
    for target, name in [(window.widgets['t55'], 'oral'), (window.details, 'result')]:
        window.scroll.ensureWidgetVisible(target)
        app.processEvents()
        window.grab().save(str(out / f'app-{name}.png'))
    window.close()

def test_tooth_error_alignment(registry, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    from PySide6.QtCore import QPoint
    from oral_registry.app import Window, configure_app
    app = QApplication.instance() or QApplication([])
    configure_app(app)
    window = Window(registry.path)
    window.show()
    app.processEvents()
    try:
        window.widgets['t51'].setText('0')
        for key in ('t55', 't54', 't51'):
            window.check_input(key)
        app.processEvents()
        positions = [window.widgets['t' + t].mapTo(window.scroll.widget(), QPoint(0, 0)).y() for t in TEETH[:10]]
        assert max(positions) - min(positions) <= 1, positions
        window.bulk_inputs['ฟันบน'].setText('0')
        window.bulk_buttons['ฟันบน'].click()
        assert not window.widgets['t55'].text()
        assert 'A' in window.bulk_errors['ฟันบน'].text()
        window.bulk_inputs['ฟันบน'].setText('a')
        window.bulk_buttons['ฟันบน'].click()
        assert window.widgets['t55'].text() == 'A'
        assert window.widgets['t51'].text() == '0'
        app.processEvents()
        window.scroll.ensureWidgetVisible(window.widgets['t55'])
        app.processEvents()
        window.grab().save('tmp/tooth-errors-fixed.png')
    finally:
        window.dirty = False
        window.close()

def test_buddhist_date_entry_and_age(registry, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    from PySide6.QtTest import QTest
    from oral_registry.app import Window, configure_app
    app = QApplication.instance() or QApplication([])
    configure_app(app)
    window = Window(registry.path)
    window.show()
    try:
        window.widgets['dob'].clear()
        QTest.keyClicks(window.widgets['dob'], '12032563')
        window.widgets['iv_date'].clear()
        QTest.keyClicks(window.widgets['iv_date'], '12092569')
        assert window.widgets['dob'].text() == '12-03-2563'
        assert window.widgets['iv_date'].text() == '12-09-2569'
        assert window.age_display.text() == '6 ปี 6 เดือน'
        assert window.age_display.isReadOnly()
        window.widgets['dob'].setText('31-02-2563')
        assert not window.check_input('dob')
        assert not window.age_display.text()
        window.widgets['dob'].setText('12-03-2020')
        assert not window.check_input('dob')
        window.widgets['dob'].setText('12-09-2570')
        assert not window.age_display.text()
        window.widgets['dob'].setText('12-03-2563')
        assert window.age_display.text() == '6 ปี 6 เดือน'
        for arch in ('ฟันบน', 'ฟันล่าง'):
            assert window.bulk_inputs[arch].text() == 'A'
        window.bulk_inputs['ฟันบน'].clear()
        window.bulk_buttons['ฟันบน'].click()
        assert all(window.widgets['t' + t].text() == 'A' for t in TEETH[:10])
        window.dirty = False
        window.table.selectRow(0)
        window.load_selected()
        assert window.widgets['dob'].text() == '06-09-2565'
        assert window.age_display.text() == '4 ปี 0 เดือน'
        app.processEvents()
        window.grab().save('tmp/date-age-preview.png')
    finally:
        window.dirty = False
        window.close()

def test_repeated_defaults(registry, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    from oral_registry.app import Window
    app = QApplication.instance() or QApplication([])
    window = Window(registry.path)
    try:
        values = dict(round_id='2', sch='โรงเรียนทดสอบ', pv='นครปฐม', hcu='รพ.สต.ทดสอบ', iv_date='12-09-2569')
        window.widgets['sch'].setText('รายการปัจจุบัน')
        window.save_settings(values)
        assert window.widgets['sch'].text() == 'รายการปัจจุบัน'
        assert not window.repeat_badges['sch'].isHidden()
        window.reset_form()
        assert all(window.widgets[k].text() == v for k, v in values.items())
        window.widgets['sch'].setText('แก้เฉพาะราย')
        window.reset_form()
        assert window.widgets['sch'].text() == values['sch']
        window.set_database(registry.path)
        assert window.repeat_values == values
        window.table.selectRow(0)
        window.load_selected()
        assert window.widgets['round_id'].text() == '1'
        assert window.widgets['sch'].text() != values['sch']
        with pytest.raises(ValueError):
            window.save_settings({'round_id': '0'})
        with pytest.raises(ValueError):
            window.save_settings({'iv_date': '31-02-2569'})
        window.save_settings({})
        window.reset_form()
        assert not window.widgets['sch'].text()
        assert window.widgets['round_id'].text() == '1'
        assert window.repeat_badges['sch'].isHidden()
    finally:
        window.dirty = False
        window.close()

def test_debris_grid_alignment(registry, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    from PySide6.QtCore import QPoint
    from oral_registry.app import Window, configure_app
    app = QApplication.instance() or QApplication([])
    configure_app(app)
    window = Window(registry.path)
    window.show()
    try:
        app.processEvents()
        panel = window.widgets['d52'].parentWidget().parentWidget()
        def positions():
            return [window.widgets['d' + t].mapTo(panel, QPoint(0, 0)) for t in FRONT]
        before = positions()
        window.widgets['d52'].setText('7')
        window.check_input('d52')
        window.check_input('d81')
        app.processEvents()
        after = positions()
        assert before == after
        assert len({p.y() for p in after[:4]}) == 1
        assert after[4].y() - after[0].y() == 32
        assert after[1].x() - after[0].x() == 52
        assert window.debris_error.text()
        assert window.field_errors['d52'].isHidden()
        assert window.debris_bulk.text() == '0'
        window.debris_bulk.setText('7')
        window.debris_fill_button.click()
        assert not window.widgets['d51'].text()
        assert 'เติมไม่ได้' in window.debris_feedback.text()
        window.debris_bulk.clear()
        window.debris_fill_button.click()
        assert window.debris_bulk.text() == '0'
        assert window.widgets['d52'].text() == '7'
        assert all(window.widgets['d' + t].text() == '0' for t in FRONT[1:])
        assert not window.widgets['vplaque'].text()
        window.scroll.ensureWidgetVisible(panel)
        app.processEvents()
        window.grab().save('tmp/debris-grid.png')
    finally:
        window.dirty = False
        window.close()

def test_auto_visible_plaque(registry, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    from PySide6.QtTest import QTest
    from oral_registry.app import Window
    app = QApplication.instance() or QApplication([])
    window = Window(registry.path)
    try:
        for t in FRONT[:-1]:
            window.widgets['d' + t].setText('3')
        assert window.widgets['vplaque'].text() == ''
        window.widgets['d72'].setText('3')
        assert window.widgets['vplaque'].text() == '1'
        for t in FRONT:
            window.widgets['d' + t].setText('0')
        assert window.widgets['vplaque'].text() == '0'
        window.widgets['vplaque'].selectAll()
        QTest.keyClicks(window.widgets['vplaque'], '1')
        window.widgets['d52'].setText('9')
        assert window.widgets['vplaque'].text() == '1'
        window.plaque_auto.click()
        assert window.widgets['vplaque'].text() == '0'
        for t in FRONT:
            window.widgets['d' + t].setText('9')
        assert window.widgets['vplaque'].text() == ''
        window.widgets['d52'].setText('2')
        assert window.widgets['vplaque'].text() == '1'
        window.widgets['d51'].clear()
        assert window.widgets['vplaque'].text() == ''
        window.reset_form()
        window.debris_fill_button.click()
        assert window.widgets['vplaque'].text() == '0'
    finally:
        window.dirty = False
        window.close()

def test_source_and_records_view(registry, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication, QFileDialog, QDialog, QTableWidget, QPushButton
    from PySide6.QtCore import Qt
    from oral_registry import app as module
    app = QApplication.instance() or QApplication([])
    module.configure_app(app)
    window = module.Window(registry.path)
    window.show()
    app.processEvents()
    try:
        saved = {}
        class Settings:
            def setValue(self, key, value):
                saved[key] = value
        monkeypatch.setattr(module, 'QSettings', lambda *args: Settings())
        monkeypatch.setattr(QFileDialog, 'getOpenFileName', lambda *args: (str(registry.path), ''))
        window.open_database()
        assert saved['data_source'] == str(registry.path)
        def inspect(dialog):
            table = dialog.findChild(QTableWidget, "recordsBody")
            assert dialog.findChild(QTableWidget, "recordsHeader").rowCount() == 3
            assert table.wordWrap()
            assert table.columnWidth(3) < 175
            assert table.columnCount() == len(registry.columns) + 2
            assert table.rowCount() == len(registry.records())
            assert table.item(0, 1).text() == '4'
            dialog.show()
            app.processEvents()
            dialog.grab().save('tmp/records-headers.png')
            dialog.close()
            return 0
        monkeypatch.setattr(QDialog, 'exec', inspect)
        window.show_records()
        buttons = window.findChildren(QPushButton)
        assert not any(b.text() == 'สร้างฐานข้อมูล' for b in buttons)
        assert all(b.cursor().shape() == Qt.CursorShape.PointingHandCursor for b in buttons if b.isVisible())
        window.table.selectRow(0)
        window.load_selected()
        assert '3 ปีขึ้นไป' in window.assessment_group.text()
        window.scroll.ensureWidgetVisible(window.details)
        app.processEvents()
        window.grab().save('tmp/assessment-ui.png')
    finally:
        window.dirty = False
        window.close()


def test_delete_registry(registry):
    original = registry.records()
    cid, round_id = original[0][1]['cid'], original[0][1]['round_id']
    backup = registry.delete(str(cid), round_id)
    assert backup.exists()
    assert len(registry.records()) == len(original) - 1
    assert Registry(backup).records()[0][1]['cid'] == cid
    with pytest.raises(ValueError):
        registry.delete(str(cid), round_id)


def test_records_edit_delete_buttons(registry, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication, QDialog, QTableWidget, QPushButton, QMessageBox
    from oral_registry.app import Window, configure_app
    app = QApplication.instance() or QApplication([])
    configure_app(app)
    window = Window(registry.path)
    try:
        def edit(dialog):
            table = dialog.findChild(QTableWidget, 'recordsBody')
            next(b for b in table.cellWidget(0, 0).findChildren(QPushButton) if b.text() == 'Edit').click()
            return 0
        monkeypatch.setattr(QDialog, 'exec', edit)
        window.show_records()
        assert window.editing is not None
        def delete(dialog):
            table = dialog.findChild(QTableWidget, 'recordsBody')
            button = next(b for b in table.cellWidget(0, 0).findChildren(QPushButton) if b.text() == 'Delete')
            monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.StandardButton.No)
            button.click()
            assert table.rowCount() == 1
            monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.StandardButton.Yes)
            button.click()
            assert table.rowCount() == 0
            return 0
        monkeypatch.setattr(QDialog, 'exec', delete)
        window.show_records()
        assert window.registry.records() == []
        assert window.editing is None
    finally:
        window.dirty = False
        window.close()
