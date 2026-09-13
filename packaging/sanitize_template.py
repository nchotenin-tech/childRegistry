from pathlib import Path
from copy import copy
import openpyxl

def sanitize(source, target):
    src=openpyxl.load_workbook(source);dst=openpyxl.Workbook();dst.remove(dst.active)
    for old in src:
        new=dst.create_sheet(old.title)
        for row in old.iter_rows(max_row=4 if old.title=='Records' else old.max_row):
            for c in row:
                if c.value is None and not c.has_style: continue
                cell=new.cell(c.row,c.column)
                for name in ('font','fill','border','alignment','protection'):
                    setattr(cell,name,copy(getattr(c,name)))
                cell.number_format=c.number_format
                if not(old.title=='Records' and c.row>=4) and not(old.title=='Entry' and c.column not in (1,2,4)):
                    cell.value=c.value
        for k,v in old.column_dimensions.items():new.column_dimensions[k].width=v.width
        for k,v in old.row_dimensions.items():
            if old.title!='Records' or k<=4:new.row_dimensions[k].height=v.height
        for merged in old.merged_cells.ranges:
            if old.title!='Records' or merged.max_row<=3:new.merge_cells(str(merged))
        new.freeze_panes=old.freeze_panes
    dst.save(target)

if __name__=='__main__':
    import sys
    sanitize(sys.argv[1],sys.argv[2])
